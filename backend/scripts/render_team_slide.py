import os
import subprocess
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches

def render_and_embed():
    edge_path = r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe"
    html_path = Path(r"c:\Users\Rishabh Bansal\Downloads\Parali setu\backend\scripts\team_slide_template.html").resolve()
    output_png = Path(r"c:\Users\Rishabh Bansal\Downloads\Parali setu\Team_Slide.png").resolve()
    slides_img = Path(r"c:\Users\Rishabh Bansal\Downloads\Parali setu\docs\slides\images\Team_Slide.png").resolve()
    output_pptx = Path(r"c:\Users\Rishabh Bansal\Downloads\Parali setu\Team_Slide.pptx").resolve()

    cmd = [
        edge_path,
        "--headless=new",
        "--disable-gpu",
        f"--screenshot={output_png}",
        "--window-size=1920,1080",
        html_path.as_uri()
    ]

    print("Rendering high-res slide via Edge headless...")
    res = subprocess.run(cmd, capture_output=True, text=True, timeout=30)
    print("Return code:", res.returncode)
    
    if output_png.exists() and output_png.stat().st_size > 0:
        print(f"Rendered successfully to {output_png} (size: {output_png.stat().st_size} bytes)")
        
        # Copy to docs/slides/images
        import shutil
        shutil.copyfile(output_png, slides_img)
        print(f"Copied to {slides_img}")

        # Embed into PPTX as full-bleed 16:9 slide
        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)
        slide_layout = prs.slide_layouts[6] # blank
        slide = prs.slides.add_slide(slide_layout)

        # Full bleed slide picture
        slide.shapes.add_picture(str(output_png), Inches(0), Inches(0), prs.slide_width, prs.slide_height)
        prs.save(str(output_pptx))
        print(f"Saved PPTX to {output_pptx}")
    else:
        print("ERROR: Screenshot was not generated!")

if __name__ == "__main__":
    render_and_embed()
